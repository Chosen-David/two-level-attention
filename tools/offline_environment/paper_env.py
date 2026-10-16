#!/usr/bin/env python3
"""Probe and recreate the paper Python environment without network or shell eval."""
import argparse
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys
import sysconfig
import venv

ROOT = Path(__file__).resolve().parents[2]
BUNDLE = ROOT / 'offline/paper'


def probe():
    return {'scope': 'machine running this probe, not a remote host',
            'os': platform.freedesktop_os_release() if sys.platform == 'linux' else {},
            'system': platform.system(), 'machine': platform.machine(),
            'python': platform.python_version(), 'implementation': platform.python_implementation(),
            'soabi': sysconfig.get_config_var('SOABI'), 'libc': platform.libc_ver(),
            'tools': {n: shutil.which(n) for n in
                      ('xelatex', 'xetex', 'inkscape', 'pdftoppm', 'pdfinfo', 'fc-match', 'nvcc', 'cmake')}}


def verify(bundle=BUNDLE):
    manifest = json.loads((bundle / 'manifest.json').read_text())
    files = manifest['files']
    if not files or len(files) != len({r['filename'] for r in files}):
        raise ValueError('empty or duplicate manifest')
    for row in files:
        name = row['filename']
        if Path(name).name != name or not name.endswith('.whl'):
            raise ValueError('unsafe wheel filename')
        path = bundle / 'wheelhouse' / name
        if path.is_symlink() or not path.is_file():
            raise ValueError(f'missing or linked wheel: {name}')
        if path.stat().st_size != row['bytes'] or hashlib.sha256(path.read_bytes()).hexdigest() != row['sha256']:
            raise ValueError(f'wheel hash/size mismatch: {name}')
    if hashlib.sha256((bundle / 'requirements.lock').read_bytes()).hexdigest() != manifest['requirements_sha256']:
        raise ValueError('requirements hash mismatch')
    return manifest


def install(destination):
    manifest = verify()
    if platform.system() != 'Linux' or platform.machine() != 'x86_64' or sys.version_info[:2] != (3, 12) or platform.python_implementation() != 'CPython':
        raise SystemExit('This bundle targets CPython 3.12 Linux x86_64 only; probe and prepare a separate bundle for other targets.')
    libc, version = platform.libc_ver()
    if libc != 'glibc' or tuple(map(int, version.split('.')[:2])) < (2, 28):
        raise SystemExit('This wheel set requires glibc >= 2.28; musl is unsupported.')
    destination = Path(destination).absolute()
    if destination.parent.resolve() != ROOT or not destination.name.startswith('.venv-') or destination.exists() or destination.is_symlink():
        raise SystemExit('Destination must be a NEW .venv-* directory directly under the repository; existing environments are never overwritten.')
    venv.EnvBuilder(with_pip=True).create(destination)
    py = destination / 'bin/python'
    env = dict(os.environ, PIP_CONFIG_FILE=os.devnull, PIP_NO_INDEX='1', PIP_DISABLE_PIP_VERSION_CHECK='1')
    subprocess.run([str(py), '-m', 'pip', '--isolated', 'install', '--no-index', '--no-cache-dir',
                    '--only-binary=:all:', '--require-hashes', '--find-links', str(BUNDLE / 'wheelhouse'),
                    '-r', str(BUNDLE / 'requirements.lock')], check=True, env=env)
    subprocess.run([str(py), '-m', 'pip', 'check'], check=True, env=env)
    subprocess.run([str(py), str(Path(__file__).resolve()), 'smoke'], check=True, env=env)
    (destination / 'paper-bundle-receipt.json').write_text(json.dumps({'manifest_sha256': hashlib.sha256((BUNDLE/'manifest.json').read_bytes()).hexdigest(), 'probe': probe()}, indent=2)+'\n')


def smoke():
    import tempfile
    import numpy as np
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from pptx import Presentation
    with tempfile.TemporaryDirectory() as td:
        fig, ax = plt.subplots(); ax.plot(np.arange(3), [1, 4, 9])
        fig.savefig(Path(td)/'smoke.pdf'); plt.close(fig)
        p = Presentation(); p.slides.add_slide(p.slide_layouts[6]); p.save(Path(td)/'smoke.pptx')
        if len(Presentation(Path(td)/'smoke.pptx').slides) != 1:
            raise RuntimeError('PPT roundtrip failed')
    print(json.dumps({'numpy': np.__version__, 'matplotlib': matplotlib.__version__, 'python-pptx': importlib.metadata.version('python-pptx'), 'synthetic_pdf_and_pptx_roundtrip': 'pass'}))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=['probe', 'verify', 'install', 'smoke'])
    parser.add_argument('--venv', default=str(ROOT/'.venv-paper'))
    args = parser.parse_args()
    if args.action == 'probe': print(json.dumps(probe(), indent=2))
    elif args.action == 'verify': print(json.dumps({'status': 'verified', 'wheels': len(verify()['files'])}))
    elif args.action == 'install': install(args.venv)
    else: smoke()

if __name__ == '__main__': main()
