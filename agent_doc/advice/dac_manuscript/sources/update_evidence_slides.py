from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pathlib import Path
R=Path(__file__).resolve().parents[1];p=R/'presentation/TwoLevel_UserEdited_20261014.pptx';r=Presentation(p)
# Bounded layout/wording refresh: preserve all other slide XML and text.
import argparse
parser = argparse.ArgumentParser(description='Synchronize archived-evidence presentation assets.')
parser.add_argument('--only', choices=['all', 'decode-ruler-q3', 'ruler32-64'], default='all')
args = parser.parse_args()
if args.only == 'ruler32-64':
 # Surgical ZIP edits preserve all other slide XML, diagrams and user material.
 from lxml import etree
 import zipfile, copy
 ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
       'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
       'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
 updates = {
 14: {2: ['RULER 32K／64K：方法配置与稠密 FullKV 的 AVG'],
      4: ['11任务 × 100例／任务／臂；32K跨机：8B β=.0625、评分器不同；32B β=.125、清单匹配。',
          '64K稀疏两臂报告同机、评分器哈希相同，但生成脚本／有效配置不同；稠密锚另次运行、评分器不同、legacy-partial。',
          '仅稀疏臂K=1024，页64／K₁=128／保护4+128／cmp4；门控与实际评分维数未核。',
          '描述性汇总；预测未重放、无配对区间，不证明因果、非劣或匹配质量—完整请求成本收益。']},
 16: {2: ['32K／64K逐任务差异解释 AVG，不另设逐项零损门槛'],
      4: ['64K：三臂MultiKey-3均为0、CWE均低于1；稀疏臂评分器哈希相同，生成脚本／有效配置不同。',
          '稠密锚另次运行、不同评分器、legacy-partial；预测未重放，无配对区间；门控与实际维数未核。']}}
 replacement={}
 with zipfile.ZipFile(p) as z:
  for number, edits in updates.items():
   name=f'ppt/slides/slide{number}.xml';xml=etree.fromstring(z.read(name))
   shapes=list(xml.find('p:cSld/p:spTree',ns))[2:]
   for index, lines in edits.items():
    body=shapes[index].find('p:txBody',ns);old=body.findall('a:p',ns);template=copy.deepcopy(old[0])
    for para in old:body.remove(para)
    for line in lines:
     para=copy.deepcopy(template);texts=para.findall('.//a:t',ns)
     texts[0].text=line
     for extra in texts[1:]:extra.text=''
     if index==4:
      for prop in para.findall('.//a:rPr',ns)+para.findall('.//a:defRPr',ns)+para.findall('.//a:endParaRPr',ns):prop.set('sz','1600' if number==14 else '1500')
     body.append(para)
   box=shapes[4].find('p:spPr/a:xfrm',ns)
   if number==14:
    box.find('a:off',ns).set('y',str(int(6.35*914400)));box.find('a:ext',ns).set('cy',str(int(1.55*914400)))
   else:
    box.find('a:off',ns).set('y',str(int(8.40*914400)));box.find('a:ext',ns).set('cy',str(int(.65*914400)))
    pic=shapes[3].find('p:spPr/a:xfrm',ns)
    pic.find('a:off',ns).set('x',str(int(3.0*914400)))
    pic.find('a:ext',ns).set('cx',str(int(10.7*914400)));pic.find('a:ext',ns).set('cy',str(int(10.7*4.5/7.1*914400)))
   picture=xml.find('.//p:pic',ns);rid=picture.find('.//a:blip',ns).get('{'+ns['r']+'}embed')
   relname=f'ppt/slides/_rels/slide{number}.xml.rels';rels=etree.fromstring(z.read(relname))
   media=f'ruler32_64_slide{number}.png'
   next(el for el in rels if el.get('Id')==rid).set('Target','../media/'+media)
   stem='ruler32k_avg_compact_zh' if number==14 else 'ruler32k_comparison_zh'
   replacement['ppt/media/'+media]=(R/'figures/ruler32k_comparison'/f'{stem}.png').read_bytes()
   replacement[name]=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
   replacement[relname]=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
  tmp=p.with_suffix('.tmp.pptx')
  with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as out:
   for item in z.infolist():out.writestr(item,replacement.pop(item.filename,z.read(item.filename)))
   for name,data in replacement.items():out.writestr(name,data)
 tmp.replace(p)
 raise SystemExit(0)
if args.only == 'decode-ruler-q3':
 from lxml import etree
 import zipfile
 import os
 ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
       'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
       'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
 replacements = {}
 with zipfile.ZipFile(p) as archive:
  for number, stem in [(12, 'decode_tpot_current_zh'), (14, 'ruler128k_current_zh')]:
   slide_path = f'ppt/slides/slide{number}.xml'
   slide = etree.fromstring(archive.read(slide_path))
   pictures = slide.findall('.//p:pic', ns)
   if len(pictures) != 1:
    raise ValueError(f'Expected exactly one evidence picture on slide {number}')
   rid = pictures[0].find('.//a:blip', ns).get('{'+ns['r']+'}embed')
   rels_path = f'ppt/slides/_rels/slide{number}.xml.rels'
   rels = etree.fromstring(archive.read(rels_path))
   relation = next(el for el in rels if el.get('Id') == rid)
   # Give the two slides their own assets; shared image users remain untouched.
   media_name = f'dac_layout_{stem}.png'
   relation.set('Target', '../media/'+media_name)
   replacements[rels_path] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
   replacements['ppt/media/'+media_name] = (R/'figures'/(stem+'.png')).read_bytes()
  path = 'ppt/slides/slide15.xml'
  slide = etree.fromstring(archive.read(path))
  edits = {
   '评分降维未降低该设置下的局部耗时': '共享子空间评分未显示一致计时优势',
   '采样协议存在未核实差异，小幅时延差值不足以定位瓶颈。':
    '相对名义页评分，两轮慢0.0237／快0.0046 ms，汇总仍较高。原始样本缺失、协议差异未解，不能认定等价或定位瓶颈。未与质量配对，未测完整请求成本。',
  }
  for old, new in edits.items():
   nodes = [el for el in slide.findall('.//a:t', ns) if el.text in (old, new)]
   if len(nodes) != 1:
    raise ValueError('Expected unique existing Q3 text: '+old)
   nodes[0].text = new
  # Give the scope note three readable lines below the chart.
  note = nodes[0]
  shape = note
  while shape.tag != '{'+ns['p']+'}sp': shape = shape.getparent()
  extent = shape.find('.//a:xfrm/a:ext', ns)
  extent.set('cy', str(int(Inches(1.08))))
  replacements[path] = etree.tostring(slide, xml_declaration=True, encoding='UTF-8', standalone=True)
  temp = p.with_suffix('.layout-tmp.pptx')
  with zipfile.ZipFile(temp, 'w', zipfile.ZIP_DEFLATED) as out:
   for item in archive.infolist(): out.writestr(item, replacements.pop(item.filename, archive.read(item.filename)))
   for name, content in replacements.items(): out.writestr(name, content)
 os.replace(temp, p)
 raise SystemExit(0)
def text(sl,x,y,w,h,value,size=22):
 s=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame;tf.word_wrap=True
 for i,line in enumerate(value.split('\n')):
  a=tf.paragraphs[0] if i==0 else tf.add_paragraph();a.text=line;a.font.name='Noto Sans CJK SC';a.font.size=Pt(size);a.font.color.rgb=RGBColor.from_string('243746');a.space_after=Pt(14)
 return s
for num,keep,title in [(3,[0,1,90],'四项观察：区域选集、紧凑坐标与历史层差异'),(12,[0,1,17],'历史动态缓存：同步模型级解码计时'),(14,[0,1,2],'保留历史128K质量负结果')]:
 sl=r.slides[num-1];orig=list(sl.shapes)
 for shape in orig:
  is_footer=shape.has_text_frame and shape.text.strip()==str(num) and shape.top>Inches(8.5)
  if (shape.top>=Inches(1.9) and not is_footer) or (num==3 and shape.shape_type==MSO_SHAPE_TYPE.PICTURE):shape._element.getparent().remove(shape._element)
 for shape in sl.shapes:
  if shape.has_text_frame and Inches(.7)<shape.top<Inches(1.8):shape.text=title
 if num==3:
  sl.shapes.add_picture(str(R/'figures/o_observation_zh.png'),Inches(2.9),Inches(1.85),height=Inches(6.55))
  text(sl,.9,8.47,14.9,.70,'(a) 页权重排序≠词元TopK；实际1989/1989/2048词元；均值／极值同32维，参考为全维；不能证明区域分工；\v(b) 43/128条下降；(c) 随机未配对；(d) 历史代理，不是当前门控或解码保证。',16)
 elif num==12:
  sl.shapes.add_picture(str(R/'figures/decode_tpot_current_zh.png'),Inches(.8),Inches(2.2),width=Inches(8))
  text(sl,9.2,2.2,6.6,6.2,'作者报告中位数 · 非完整请求E2E\nQwen3-8B，固定128输出词元\n127个首输出后间隔，逐步同步\n历史稀疏配置 α=.125 / β=.375 / γ=.625\n不同于当前4/128候选；\v运行路径未独立核验\n32K/64K/128K稀疏／稠密中位数比：\v3.38 / 2.46 / 1.78\n三档比值减小，但稀疏TPOT仍上升且高于稠密；汇总不能定位因果瓶颈。',23)
 else:
  sl.shapes.add_picture(str(R/'figures/ruler128k_current_zh.png'),Inches(.8),Inches(2.3),width=Inches(8))
  text(sl,9.2,2.2,6.6,6.2,'历史128K得分：46.43 / 23.98\n11任务；每任务每臂100例\n报告 β=.375，K1=128，保护128/128\n中间预算768；页长64，γ=.625\n按已检公式：名义远区词元配额=0\n代码／配置推论；实际执行未核验\n不能据此认定得分差的唯一原因；不验证当前4/128候选。',23)
# Terminology and attribution only; retain the existing17-slide structure.
for slide in r.slides:
 for shape in slide.shapes:
  if not shape.has_text_frame:continue
  for para in shape.text_frame.paragraphs:
   for run in para.runs:
    run.text=run.text.replace('远区概率质量','远区注意力权重总和').replace('远区 概率求和 → 均值','标定区注意力求和 → 均值').replace('组共享的两级候选选择','页到词元的级联稀疏选择').replace('全维页评分、校准精筛通道与异步执行','相关工作的级联选择先例')
# Keep the statistic label within its existing gate node.
label=r.slides[8].shapes[84]
label.left=Inches(3.88); label.top=Inches(5.06)
label.width=Inches(1.04); label.height=Inches(.35)
label.text_frame.margin_left=0; label.text_frame.margin_right=0
label.text_frame.margin_top=0; label.text_frame.margin_bottom=0
label.text_frame.paragraphs[0].text='远区注意力\v代理统计'
for para in label.text_frame.paragraphs:
 for run in para.runs:
  run.font.size=Pt(10); run.font.name='Noto Sans CJK SC'; run.font.color.rgb=RGBColor.from_string('7B6598')

# Full-manuscript narrative revision: preserve17slides and existing diagrams.
updates = {
 (9,2): '门控默认关闭；标定区域与跳过页边界不同，统计量是代理（补充S5）',
 (7,2): '16完整旋转对，两级共用；L2用浮点16级量化–反量化，最终为原维注意力',
 (8,2): '4/128候选充足示例；摘要按完整物理页计算，远区可重复选中至多4个前缀词元',
 (8,158): '近区L2∩P=∅；|Pₜ|=132示例',
 (2,1): '固定词元预算下：候选访问与选择成本',
 (2,3): '均值与极值刻画不同评分目标，页注意力和另作参照。\n全局TopK不提供区域容量保证。',
 (2,5): '共享32坐标映射与区域页配额协同选择，\n最终在选集上使用原128维 Q / K / V。',
 (2,6): '观察与边界',
 (2,7): '区域预留改变远区覆盖，43/128条观测下降；历史稀疏解码更慢，当前4/128尚缺匹配质量—成本结果。',
 (4,9): '级联、坐标选择与量化评分',
 (4,10): '逐查询概率聚合；原维KV最终注意力',
 (5,1): '两项核心机制，一个可选扩展',
 (5,5): '共享坐标映射，不要求共享物理缓冲区',
 (11,1): 'Q1：缩小评分宽度能否保留任务质量？',
 (13,1): '区域访问与门控：证据支持到哪里？',
 (13,6): '五任务均值；相对名义全维页评分\n低0.322分，相对名义全维词元评分\n低0.078分；配置未确认匹配，\n不能隔离阶段效应。',
 (17,1): '评分表示、区域访问与执行成本的取舍',
 (17,12): '保留远区页，不保证留有远区词元配额\n当前4/128尚缺匹配质量—成本验证',
 (17,13): '门控默认关闭；历史任务存在损失\n预填充统计不保证解码需求',
}
for (num, idx), value in updates.items():
 tf=r.slides[num-1].shapes[idx].text_frame
 runs=[run for para in tf.paragraphs for run in para.runs]
 if runs:
  runs[0].text=value
  for run in runs[1:]: run.text=''
 else: tf.paragraphs[0].text=value

r.save(p)
