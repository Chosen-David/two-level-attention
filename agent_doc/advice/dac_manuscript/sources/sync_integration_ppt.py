from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree
R=Path(__file__).resolve().parents[1]
P=R/'presentation/TwoLevel_UserEdited_20261014.pptx'
p=Presentation(P)
assert len(p.slides)==17
def replace_text(shape,text):
 tf=shape.text_frame
 first=tf.paragraphs[0]
 # Preserve existing font/paragraph properties while replacing visible content.
 if first.runs:
  first.runs[0].text=text
  for r in list(first.runs)[1:]:r._r.getparent().remove(r._r)
 else:first.text=text
 for par in list(tf.paragraphs)[1:]:par._p.getparent().remove(par._p)
def replace_picture(slide,old,path):
 new=slide.shapes.add_picture(str(path),old.left,old.top,old.width,old.height)
 old._element.getparent().remove(old._element)
 return new
replace_text(p.slides[1].shapes[7],'极值在两区的漏失均较低；混合配额诊断未胜全局极值（.135 / .121）。历史稀疏解码更慢，当前设计尚缺匹配质量—成本结果。')
s=p.slides[2]
replace_picture(s,next(sh for sh in s.shapes if sh.shape_type==13),R/'figures/o_observation_zh.png')
replace_text(s.shapes[4] if s.shapes[4].has_text_frame else s.shapes[3], '(a) 均值／极值同32维，参考为全维；实际词元数不同；(b) 各区漏失 ÷ 每查询参考集总质量，再取均值；\n极值两区均较低；23种拆分最佳 .135，仍差于全局极值 .121；(c) 随机未配对；(d) 历史代理，非解码保证。')
replace_text(p.slides[4].shapes[9],'显式近区页／词元配额，远区按差值\n最终使用原128维 Q / K / V')
s=p.slides[5]
replace_text(s.shapes[2],'分区参考路径：边界保送开启，门控／候选收集关闭；示例数字不是默认配置；归档评测不验证新路径')
for sh in list(s.shapes)[4:]:sh._element.getparent().remove(sh._element)
scene=json.loads((R/'figures/architecture_domains/boundary_scene_zh.json').read_text())
scale=6.6/600
ox,oy=(16.6666666667-1100*scale)/2,2.04
def X(x):return Inches(ox+x*scale)
def Y(y):return Inches(oy+y*scale)
def color(v):return RGBColor.from_string(v.lstrip('#'))
for o in scene['objects']:
 if o['type']=='rect':
  sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,X(o['x']),Y(o['y']),Inches(o['w']*scale),Inches(o['h']*scale))
  if o['fill']=='none':sh.fill.background()
  else:sh.fill.solid();sh.fill.fore_color.rgb=color(o['fill'])
  sh.line.color.rgb=color(o['stroke']);sh.line.width=Pt(o['sw']*scale*72)
 elif o['type']=='line':
  sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,X(o['x1']),Y(o['y1']),X(o['x2']),Y(o['y2']))
  sh.line.color.rgb=color(o['color']);sh.line.width=Pt(1)
  if o['arrow']:
   ln=sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
   etree.SubElement(ln,'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd',type='triangle',w='sm',len='sm')
 else:
  width=1100 # Centered text box retains exact center; no word-wrap.
  sh=s.shapes.add_textbox(X(o['x']-width/2),Y(o['y']-o['size']*1.05),Inches(width*scale),Inches(o['size']*1.5*scale))
  tf=sh.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
  par=tf.paragraphs[0];par.alignment=PP_ALIGN.CENTER
  run=par.add_run();run.text=o['text'];run.font.name='Noto Sans CJK SC';run.font.size=Pt(o['size']*scale*72);run.font.color.rgb=color(o['color'])
replace_text(p.slides[6].shapes[2],'16完整旋转对，两级共用；图示QDQ为历史量化变体，主文bf16配置不以该量化为必要条件；最终为原维注意力')
replace_text(p.slides[7].shapes[2],'本页展示另一组近区配额(16,640)，非主表(8,320)；边界页额外获L2资格，非整页最终必选（见第6页）')
replace_text(p.slides[8].shapes[2],'门控默认关闭；边界页仍可能进入评分，不能笼统声称省去全部远区点积；统计量非解码保证')
for idx,fig in [(13,'ruler32k_avg_compact_zh.png'),(15,'ruler32k_comparison_zh.png')]:
 slide=p.slides[idx]
 replace_picture(slide,next(sh for sh in slide.shapes if sh.shape_type==13),R/'figures/ruler32k_comparison'/fig)
 if idx==13:
  replace_text(slide.shapes[2],'Qwen3-8B：RULER 32K／64K 与稠密 FullKV 的 AVG')
  replace_text(slide.shapes[3],'11任务 × 100例／任务／臂；32K比较跨机器、评分器不同。\n64K稀疏两臂报告同机、评分器相同，生成脚本／有效配置不同；稠密锚另次运行、不同评分器、legacy-partial。\n仅稀疏臂 K=1024、4/128保护；门控与评分维数未核，不证非劣或匹配质量—成本收益。')
for sh in p.slides[8].shapes:
 if sh.has_text_frame and 'OFF：跳过远区点积' in sh.text:
  replace_text(sh,sh.text.replace('OFF：跳过远区点积','OFF：移除远区竞争'))
 if sh.has_text_frame and 'L1 k=0；L2跳过' in sh.text:
  replace_text(sh,'L1 k=0；远区无L2配额')
assert len(p.slides)==17
p.save(P)
